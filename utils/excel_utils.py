from openpyxl import load_workbook

def get_products_from_excel(file_path):
    workbook = load_workbook(file_path)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    products = []

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
        product = dict(zip(headers, row))
        if product.get("Name Of Item") is not None:
            products.append(product)

    print("TOTAL PRODUCTS:", len(products))
    workbook.close()
    return products


""" 
zip() : is used here to match each Excel column name (headers) with the value in that row (row).

product = dict(zip(headers, row))
Example :
{
    "Name Of Item": "Pearl Necklace",
    "Price": 5000,
    "Unit": "Piece"
}
"""